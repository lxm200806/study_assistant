"""Parse official 4500 Key English Words Excel lists into a catalog JSON."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
MATERIALS = ROOT / "materials" / "reference material" / "4500 Key English Words"
OUT = ROOT / "data" / "sources" / "kew" / "kew4500-units.json"

BOOK_META = {
    1: {"code": "kew4500-1", "name": "4500高频词 1", "level": "A2+", "targetWordCount": 800},
    2: {"code": "kew4500-2", "name": "4500高频词 2", "level": "B1", "targetWordCount": 800},
    3: {"code": "kew4500-3", "name": "4500高频词 3", "level": "B1+", "targetWordCount": 800},
    4: {"code": "kew4500-4", "name": "4500高频词 4", "level": "B2", "targetWordCount": 800},
}


def normalize_word(raw: str) -> str:
    return re.sub(r"\s+", " ", str(raw or "").strip().lower())


def normalize_pos(raw: str) -> str:
    text = str(raw or "").strip().lower()
    text = text.replace(" ", "")
    text = text.replace(".", "")
    return text


def parse_book(n: int) -> dict:
    matches = list((MATERIALS / f"4500 Key English Words {n}").glob("*Word Lists*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"missing Word Lists xlsx for book {n}")
    path = matches[0]
    wb = load_workbook(path, read_only=True, data_only=True)
    units: dict[str, dict] = {}
    order: list[str] = []
    current_id = None
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            a = row[0] if row else None
            b = row[1] if row and len(row) > 1 else None
            c = row[2] if row and len(row) > 2 else None
            d = row[3] if row and len(row) > 3 else None
            if isinstance(a, str) and "Unit" in a:
                m = re.search(r"Unit\s+(\d+)", a, re.I)
                if not m:
                    continue
                current_id = f"u{int(m.group(1))}"
                if current_id not in units:
                    units[current_id] = {
                        "id": current_id,
                        "title": f"Unit {int(m.group(1))}",
                        "entries": [],
                    }
                    order.append(current_id)
            elif isinstance(a, int) and current_id and isinstance(b, str):
                word = normalize_word(b)
                if not word or word == "word":
                    continue
                units[current_id]["entries"].append(
                    {
                        "word": word,
                        "pos": str(c or "").strip(),
                        "englishMeaning": re.sub(r"\s+", " ", str(d or "").strip()),
                    }
                )
    wb.close()
    meta = BOOK_META[n]
    unit_list = []
    for uid in order:
        unit = units[uid]
        unit["words"] = [entry["word"] for entry in unit["entries"]]
        unit_list.append(unit)
    return {
        "code": meta["code"],
        "name": meta["name"],
        "description": (
            f"Seed Learning《4500 Key English Words {n}》。来自官方 Word Lists，"
            "每单元 20 个 New Words，不含课文句子用词。"
        ),
        "level": meta["level"],
        "targetWordCount": meta["targetWordCount"],
        "units": unit_list,
    }


def main() -> None:
    books = [parse_book(n) for n in range(1, 5)]
    catalog = {
        "series": "4500 Key English Words",
        "publisher": "Seed Learning",
        "author": "Paul Nation",
        "source": "Official Word Lists xlsx — unit target words only (not passage running words)",
        "books": books,
    }

    by_word: dict[str, list] = defaultdict(list)
    for book in books:
        for unit in book["units"]:
            for entry in unit["entries"]:
                by_word[entry["word"]].append(
                    {
                        "book": book["code"],
                        "unit": unit["id"],
                        "pos": entry["pos"],
                        "en": entry["englishMeaning"],
                    }
                )

    multi = {w: v for w, v in by_word.items() if len(v) > 1}
    diff_pos = []
    for word, rows in multi.items():
        poss = {normalize_pos(r["pos"]) for r in rows}
        if len(poss) > 1:
            diff_pos.append({"word": word, "rows": rows})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"wrote {OUT}")
    for book in books:
        counts = [len(u["entries"]) for u in book["units"]]
        print(
            f"  {book['code']}: {len(book['units'])} units, "
            f"{sum(counts)} words, per-unit {min(counts)}-{max(counts)}"
        )
    print(f"  repeated spellings: {len(multi)}")
    print(f"  different POS: {len(diff_pos)}")
    for item in diff_pos[:30]:
        print("   ", item["word"], [(r["book"], r["unit"], r["pos"]) for r in item["rows"]])


if __name__ == "__main__":
    main()
